# -*- coding: utf-8 -*-
"""Реестр каталога WED Secrets (Google-таблица модератора).
Запуск из корня: node scripts/filters-dump.mjs && python scripts/registry-sheet.py
Строка 1 — группа, строка 2 — ключ для скрипта (серый), строка 3 — подпись. Данные с 4-й.
Закреплены столбцы «Статус» и «Название»."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
import json, glob

OUT = "WED Secrets — реестр каталога.xlsx"  # запуск из корня проекта

FONT = "Arial"
F_KEY = Font(name=FONT, size=8, color="999999")
F_LABEL = Font(name=FONT, size=10, bold=True)
F_GROUP = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_BODY = Font(name=FONT, size=10)
F_TITLE = Font(name=FONT, size=14, bold=True)
FILL_GROUP = PatternFill("solid", fgColor="2B2B2B")
FILL_SITE = PatternFill("solid", fgColor="7A7A7A")     # группа «заполняет сайт»
FILL_LABEL = PatternFill("solid", fgColor="EFEFEF")
FILL_REQ = PatternFill("solid", fgColor="FFF2CC")
FILL_DIM = PatternFill("solid", fgColor="E0E0E0")
THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")

SITE = "Заполняет сайт"
STATUS = ["черновик", "в работе", "на проверке", "опубликовано", "архив"]

REF = {
    "Статус": STATUS,
    "Тип площадки": ["ресторан", "банкетный зал", "шатёр", "веранда", "ресторан при отеле",
                     "загородный клуб", "лофт", "оранжерея", "усадьба", "яхт-клуб", "другое"],
    "Округ / направление": ["ЦАО", "САО", "СВАО", "ВАО", "ЮВАО", "ЮАО", "ЮЗАО", "ЗАО", "СЗАО", "Подмосковье"],
    # основные направления — список должен влезать в 250 знаков встроенного списка xlsx
    "Шоссе": ["в Москве", "Ленинградское", "Новорижское", "Волоколамское", "Пятницкое", "Рублёвское",
              "Минское", "Можайское", "Киевское", "Калужское", "Варшавское", "Симферопольское",
              "Каширское", "Новорязанское", "Егорьевское", "Горьковское", "Щёлковское", "Ярославское", "Дмитровское"],
    "Кто": ["мужчина", "женщина", "пара", "команда"],
    "Возраст": ["до 25", "25–35", "35–50", "50+"],
    "Округа подрядчика": ["все", "ЦАО", "САО", "СВАО", "ВАО", "ЮВАО", "ЮАО", "ЮЗАО", "ЗАО", "СЗАО"],
    "Языки": ["русский", "английский", "другие"],
    "Стиль ведения": ["классический", "современный", "интерактивный", "юмористический", "интеллигентный",
                      "с элементами шоу", "без пошлости", "на двух языках"],
    "Форматы свадеб": ["камерные (до 30)", "средние (30–80)", "масштабные (80+)", "выездные церемонии",
                       "гендер-пати", "юбилеи", "помолвки"],
    "Стиль съёмки": ["репортаж", "классика", "минимализм", "чёрно-белый", "художественный", "постановочный", "лайфстайл"],
    "Тип съёмки": ["полный день", "частичная съёмка", "лав-стори", "экспресс"],
    "Сроки выдачи": ["1–2 недели", "2–4 недели", "1–3 месяца", "3–6 месяцев", "более 6 месяцев"],
    "Формат выдачи": ["онлайн-галерея", "флешка", "фотокнига"],
    "Ценовой сегмент": ["бюджетный", "средний", "премиум"],
    "Стиль декора": ["классика", "лофт", "рустик", "бохо", "минимализм", "гламур", "эко", "восточный"],
    "Тип услуги декора": ["под ключ", "отдельные услуги", "аренда"],
    "Типы площадок (декор)": ["лофт", "банкетный зал", "усадьба", "терраса", "загородный клуб"],
    "Предоплата": ["до 20%", "20–50%", "50–100%"],
}

# (key, label, group, kind, extra) — kind: text|long|int|num|yes|url|list:<ref>|multi:<ref>
VENUE_COLS = [
    ("status", "Статус", "Основное", "list:Статус", {"req": True}),
    ("name", "Название", "Основное", "text", {"req": True}),
    ("type", "Тип площадки", "Основное", "list:Тип площадки", {"req": True}),
    ("kicker", "Бирка: тип · место", "Основное", "text", {"note": "Одна строка под именем: «Оранжерея · Быково»"}),
    ("lead", "Короткое описание", "Основное", "long", {}),
    ("city", "Населённый пункт", "Где", "text", {"req": True}),
    ("district", "Округ / направление", "Где", "list:Округ / направление", {"req": True}),
    ("highway", "Шоссе", "Где", "list:Шоссе", {}),
    ("kmFromMkad", "Км от МКАД", "Где", "num", {}),
    ("mapsUrl", "Яндекс.Карты (ссылка)", "Где", "url", {"req": True, "note": "Адрес и координаты вытянем отсюда сами"}),
    ("capacityBanquet", "Банкет, гостей до", "Вместимость", "int", {"req": True}),
    ("capacityBuffet", "Фуршет, гостей до", "Вместимость", "int", {}),
    ("hallsCount", "Залов", "Вместимость", "int", {"note": "Больше одного — залы на вкладке «Залы»"}),
    ("checkFrom", "Средний чек от, ₽/чел", "Деньги", "int", {"note": "Пусто = цена по запросу"}),
    ("rentFrom", "Аренда от, ₽", "Деньги", "int", {}),
    ("depositFrom", "Мин. заказ / депозит от, ₽", "Деньги", "int", {}),
    ("serviceFee", "Сервисный сбор, %", "Деньги", "num", {}),
    ("f_ownAlcohol", "Свой алкоголь", "Особенности", "yes", {}),
    ("f_water", "У воды", "Особенности", "yes", {}),
    ("f_ceremony", "Выездная регистрация", "Особенности", "yes", {}),
    ("f_countryside", "За городом", "Особенности", "yes", {}),
    ("f_exclusive", "Под закрытие", "Особенности", "yes", {}),
    ("f_lodging", "Проживание", "Особенности", "yes", {}),
    ("f_ownKitchen", "Своя кухня", "Особенности", "yes", {}),
    ("phone", "Телефон", "Контакты", "text", {"req": True}),
    ("email", "Email", "Контакты", "text", {}),
    ("site", "Сайт", "Контакты", "url", {}),
    ("social", "Соцсети (ссылки)", "Контакты", "long", {}),
    ("materials", "Фото и видео (ссылка на диск)", "Материалы", "url", {"req": True}),
    ("priceDoc", "Прайс / меню (ссылка)", "Материалы", "url", {}),
    ("description", "Текст от площадки (как прислали)", "Материалы", "long", {}),
    ("slug", "Slug", SITE, "text", {}),
    ("siteUrl", "Страница на сайте", SITE, "url", {}),
    ("address", "Адрес", SITE, "long", {}),
    ("geo", "Координаты", SITE, "text", {}),
    ("rating", "Рейтинг Я.Карт", SITE, "num", {}),
    ("reviewsCount", "Отзывов", SITE, "int", {}),
    ("notes", "Комментарий", SITE, "long", {}),
]

HALL_COLS = [
    ("venueName", "Площадка", "Зал", "text", {"req": True}),
    ("name", "Зал", "Зал", "text", {"req": True}),
    ("kicker", "Тип зала", "Зал", "text", {"note": "Стеклянная оранжерея, шатёр, банкетный зал"}),
    ("areaM2", "Площадь, м²", "Зал", "num", {}),
    ("banquet", "Банкет, гостей до", "Зал", "int", {"req": True}),
    ("buffet", "Фуршет, гостей до", "Зал", "int", {}),
    ("rentWeekday", "Аренда Пн–Чт, ₽", "Аренда", "int", {}),
    ("rentFri", "Аренда Пт, ₽", "Аренда", "int", {}),
    ("rentSat", "Аренда Сб, ₽", "Аренда", "int", {}),
    ("rentSun", "Аренда Вс, ₽", "Аренда", "int", {}),
    ("depositWeekday", "Мин. заказ Пн–Чт, ₽", "Мин. заказ", "int", {}),
    ("depositFri", "Мин. заказ Пт, ₽", "Мин. заказ", "int", {}),
    ("depositSat", "Мин. заказ Сб, ₽", "Мин. заказ", "int", {}),
    ("depositSun", "Мин. заказ Вс, ₽", "Мин. заказ", "int", {}),
    ("included", "Что входит (текстом)", "Описание", "long", {}),
    ("materials", "Фото зала (ссылка)", "Описание", "url", {}),
    ("hallId", "Id зала", SITE, "text", {}),
]

SPEC_COMMON = [
    ("status", "Статус", "Основное", "list:Статус", {"req": True}),
    ("name", "Имя / название", "Основное", "text", {"req": True}),
    ("who", "Кто", "Основное", "list:Кто", {}),
    ("tagline", "Бирка: одна строка под именем", "Основное", "text", {"req": True}),
    ("bio", "О себе (текст как прислали)", "Основное", "long", {"req": True}),
    ("alsoCategories", "Также в разделах", "Основное", "text", {"note": "Пара «фото + видео» — в столбце второй раздел, визитка одна"}),
    ("districts", "Округа Москвы", "Где", "multi:Округа подрядчика", {"req": True, "note": "«все» или несколько округов. В Google Таблице: Данные → Проверка данных → включить «Разрешить несколько вариантов» — будут плашки с выбором"}),
    ("mo", "Московская область", "Где", "yes", {"req": True}),
    ("outsideMo", "Выезд за МО", "Где", "yes", {}),
    ("priceFrom", "Цена от, ₽ (число для фильтра)", "Цена", "int", {"note": "Пусто = по запросу"}),
    ("priceText", "Цена как прислали", "Цена", "long", {"note": "«8 000 ₽/час, свадьба от 40 000», «за проект по смете» — своими словами"}),
    ("experienceYears", "Опыт, лет", "Опыт", "int", {}),
    ("age", "Возраст", "Опыт", "list:Возраст", {}),
    ("languages", "Языки", "Опыт", "multi:Языки", {}),
]

SPEC_CONTACTS = [
    ("phone", "Телефон", "Контакты", "text", {"req": True}),
    ("email", "Email", "Контакты", "text", {}),
    ("site", "Сайт", "Контакты", "url", {}),
    ("social", "Соцсети (ссылки)", "Контакты", "long", {}),
    ("materials", "Портфолио (ссылка на диск)", "Материалы", "url", {"req": True}),
    ("avatar", "Портрет (ссылка на файл)", "Материалы", "url", {"note": "Отдельный файл: лицо крупно"}),
    ("slug", "Slug", SITE, "text", {}),
    ("siteUrl", "Страница на сайте", SITE, "url", {}),
    ("notes", "Комментарий", SITE, "long", {}),
]

G = "Фильтры раздела"
GENERIC = [("specifics", "Особенности услуги (текстом)", G, "long", {"note": "Формат, состав, что входит — как прислали; фильтров у раздела на сайте пока нет"})]

# Фильтры категорий — из src/data/specialist-filters.ts через scripts/filters-dump.mjs.
# Эти группы уже покрыты общим блоком и в столбцы раздела не идут:
SKIP_KEYS = {"city", "gender", "age", "experience", "languages"}
PRICE_KEYS = {"price", "priceHour", "pricePerson", "priceKg"}  # бакеты цены — считаются из «Цена от, ₽»
FILTERS = json.load(open("research/filters/filters.json", encoding="utf-8"))
REF_BY_OPTIONS = {}

def ref_name(label, options):
    """один справочник на одинаковый набор вариантов; разные наборы с одной подписью — с суффиксом"""
    key = tuple(o["label"] for o in options)
    if key in REF_BY_OPTIONS:
        return REF_BY_OPTIONS[key]
    name = label
    n = 2
    while name in REF:
        name = f"{label} ({n})"; n += 1
    REF[name] = list(key); REF_BY_OPTIONS[key] = name
    return name

def category_cols(slug):
    groups = FILTERS.get(slug)
    if not groups:
        return GENERIC, None
    cols, price_label = [], None
    for g in groups:
        if g["key"] in SKIP_KEYS:
            continue
        if g["key"] in PRICE_KEYS:
            price_label = g["label"]; continue
        labels = [o["label"] for o in g["options"]]
        if labels == ["Да", "Нет"]:
            cols.append((g["key"], g["label"], G, "yes", {}))
        elif g["type"] == "multi":
            cols.append((g["key"], g["label"], G, "multi:" + ref_name(g["label"], g["options"]), {}))
        else:
            cols.append((g["key"], g["label"], G, "list:" + ref_name(g["label"], g["options"]), {}))
    return cols, price_label

def label_of(slug, key, value):
    for g in FILTERS.get(slug, []):
        if g["key"] == key:
            for o in g["options"]:
                if o["value"] == value:
                    return o["label"]
    return value

def spec_row(s):
    """строка таблицы из JSON карточки (src/data/specialists) — значения → подписи"""
    slug = s["categorySlug"]; f = s.get("filters", {})
    row = dict(status="опубликовано", name=s["name"], tagline=s.get("tagline"), bio=s.get("bio"),
               priceFrom=s.get("priceFrom"), priceText=s.get("priceNote"), experienceYears=s.get("experienceYears"), age=s.get("age"),
               slug=s["slug"], siteUrl=f"https://wed-secrets.ru/moskva/podryadchiki/{slug}/{s['slug']}/",
               materials="в репозитории", avatar="есть" if s.get("avatar") else None)
    who = {"м": "мужчина", "ж": "женщина", "pair": "пара"}.get(s.get("gender") or f.get("gender"))
    if who: row["who"] = who
    if s.get("alsoCategories"):
        row["alsoCategories"] = ", ".join(dict(CATEGORIES)[c] for c in s["alsoCategories"])
    city = f.get("city") or s.get("cities") or []
    okr = [c.upper() for c in city if c not in ("mo", "вне-мо", "moskva", "other")]
    okr = ["ЮВАО" if c == "ЮВО" else c for c in okr]
    row["districts"] = "все" if len(okr) >= 9 or "moskva" in city else ", ".join(okr)
    row["mo"] = "да" if "mo" in city else None
    row["outsideMo"] = "да" if "вне-мо" in city else ("да" if f.get("travel") == "yes" else None)
    langs = f.get("languages") or s.get("languages")
    if langs: row["languages"] = ", ".join({"ru": "русский", "en": "английский"}.get(l, l) for l in langs)
    for key, val in f.items():
        if key in SKIP_KEYS or key in PRICE_KEYS: continue
        if isinstance(val, list):
            row[key] = ", ".join(label_of(slug, key, v) for v in val)
        elif val in ("yes", "no"):
            row[key] = "да" if val == "yes" else "нет"
        else:
            row[key] = label_of(slug, key, val)
    return row


CATEGORIES = [
    ("organizatory", "Организаторы"), ("koordinatory", "Координаторы"), ("vedushchie", "Ведущие"),
    ("dekoratory", "Декораторы"), ("fotografy", "Фотографы"), ("videografy", "Видеографы"),
    ("rils-meikery", "Reels-мейкеры"), ("keitering", "Кейтеринг"), ("konditery", "Кондитеры"),
    ("stilisty", "Стилисты и визажисты"), ("dj", "Диджеи"), ("kaver-gruppy", "Кавер-группы"),
    ("vokalisty", "Вокалисты"), ("muzykanty", "Музыканты"), ("speceffekty", "Спецэффекты"),
    ("arenda-zvuka", "Аренда звука"), ("arenda-sveta", "Аренда светомузыки"), ("horeografy", "Хореографы"),
    ("animatory", "Аниматоры"), ("shou", "Шоу"), ("fokusniki", "Фокусники и иллюзионисты"), ("avto", "Авто и трансфер"),
]

# ---------- уже заведённое ----------
VENUES = [
    dict(status="опубликовано", name="Ривер Лофт", type="лофт", kicker="Площадка · Подольск, Дубровицы",
         lead="Панорамный лофт 200 м² у слияния Пахры и Десны, работает только под закрытие.",
         city="Подольск, Дубровицы", district="Подмосковье", highway="Варшавское", kmFromMkad=26, mapsUrl="https://yandex.ru/maps/-/CTtYZIZC",
         capacityBanquet=100, capacityBuffet=150, hallsCount=1, checkFrom=5000, rentFrom=30000, depositFrom=200000, serviceFee=10,
         f_ownAlcohol="да", f_water="да", f_ceremony="нет", f_countryside="да", f_exclusive="да", f_lodging="нет", f_ownKitchen="да",
         phone="+7 925 859 22 25", email="info@river-loft.ru", site="https://river-loft.ru/", social="https://t.me/RiverLoft_podolsk\nhttps://vk.com/river_loft",
         materials="фото с сайта (в репозитории)", priceDoc="https://river-loft.ru/menu.html",
         slug="river-loft", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/river-loft/", address="Московская область, г.о. Подольск, посёлок Дубровицы, д. 38", geo="55.440658, 37.491289", rating=5, reviewsCount=151),
    dict(status="опубликовано", name="Лесная Роса", type="оранжерея", kicker="Оранжерея · Быково",
         lead="Полностью стеклянная оранжерея в парке бывшей усадьбы, 14 км от МКАД.",
         city="Подольск, Быково", district="Подмосковье", highway="Варшавское", kmFromMkad=14, mapsUrl="https://yandex.ru/maps/-/CTtYbC4x",
         capacityBanquet=70, capacityBuffet=100, hallsCount=1, checkFrom=9000, rentFrom=280000, depositFrom=350000, serviceFee=15,
         f_ownAlcohol="да", f_water="да", f_ceremony="да", f_countryside="да", f_exclusive="да", f_lodging="нет", f_ownKitchen="да",
         phone="+7 993 594 71 24", email="info@forestdew.ru", site="https://forestdew.ru/", social="https://t.me/ForestDewMsk\nhttps://vk.com/forestdewwed",
         materials="https://disk.yandex.ru/d/wm5FBlRvEk3J1w\nhttps://disk.yandex.ru/d/RcDBGp2nNzp6h",
         slug="forest-dew", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/forest-dew/", address="Московская область, Подольский район, посёлок Быково, ул. Луговая, 17", geo="55.468076, 37.599354", rating=5, reviewsCount=279),
    dict(status="опубликовано", name="Из-за любви", type="оранжерея", kicker="Оранжерея · Быково",
         city="Подольск, Быково", district="Подмосковье", highway="Варшавское", kmFromMkad=15, mapsUrl="https://yandex.ru/maps/-/CTtYvUp0",
         capacityBanquet=60, capacityBuffet=80, hallsCount=1, checkFrom=9000, serviceFee=15,
         f_ownAlcohol="да", f_water="да", f_ceremony="да", f_countryside="да", f_exclusive="да", f_lodging="нет", f_ownKitchen="да",
         phone="+7 933 399 44 95", email="info@duetolove.ru", site="https://duetolove.ru/", social="https://t.me/DueToLoveRu\nhttps://vk.com/duetolove",
         materials="https://disk.yandex.ru/d/8s8pl8ZGNpXN2g\nhttps://disk.yandex.ru/d/aRZA5kjQ__znXw",
         slug="due-to-love", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/due-to-love/", address="Московская область, Подольский район, посёлок Быково, ул. Луговая, 17", geo="55.468015, 37.599379", rating=5, reviewsCount=141),
    dict(status="опубликовано", name="Spark Hall", type="банкетный зал", kicker="Банкетный зал · Новая Рига",
         lead="Зал 188 м² с потолками 6 м в сосновом парке, рядом открытая площадка 500 м².",
         city="Красногорск", district="Подмосковье", highway="Новорижское", mapsUrl="https://yandex.ru/maps/-/CTt4vZih",
         capacityBanquet=120, capacityBuffet=300, hallsCount=1,
         f_ceremony="да", f_countryside="да", f_exclusive="да",
         phone="+7 917 508-68-88", site="https://topeventhall.ru/", social="https://t.me/izoldazi\nhttps://www.instagram.com/topeventhall",
         materials="https://disk.yandex.ru/d/_-ASNPPPDgOZZQ",
         slug="spark-hall", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/spark-hall/", address="Московская область, г.о. Красногорск", geo="55.803208, 37.282198", rating=4.5, reviewsCount=16),
    dict(status="опубликовано", name="LЁD", type="оранжерея", kicker="Оранжерея · Новая Рига",
         city="Красногорск", district="Подмосковье", highway="Новорижское",
         capacityBanquet=60, hallsCount=1, f_countryside="да", f_exclusive="да",
         phone="+7 917 508-68-88", site="https://topeventhall.ru/", social="https://t.me/izoldazi",
         materials="https://disk.yandex.ru/d/VZSScWIAvbj_N",
         slug="led", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/led/", address="Московская область, г.о. Красногорск", geo="55.802866, 37.281575", rating=4.5, reviewsCount=16),
    dict(status="опубликовано", name="Отражение", type="оранжерея", kicker="Оранжерея · Можайский район",
         city="Можайский район, Малое Новосурино", district="Подмосковье", highway="Минское", mapsUrl="https://yandex.ru/maps/-/CTtaA4nb",
         capacityBanquet=60, hallsCount=1, checkFrom=8000,
         f_water="да", f_ceremony="да", f_countryside="да", f_lodging="да", f_ownKitchen="да",
         phone="+7 995 114 32 65", email="naostrovahhotel@gmail.com", site="https://oranzhereya-otrazhenie.zhivopisnyy-ostrov.ru/",
         materials="фото с сайта (заменить, когда пришлют)",
         slug="otrazhenie", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/otrazhenie/", address="Московская область, Можайский район, д. Малое Новосурино", geo="55.471452, 35.980837", rating=4.4, reviewsCount=10),
    dict(status="опубликовано", name="Fish Point", type="банкетный зал", kicker="Банкетный зал · Подольск",
         city="Подольск, Бережки", district="Подмосковье", highway="Симферопольское", mapsUrl="https://yandex.ru/maps/-/CTtaAPOP",
         capacityBanquet=200, capacityBuffet=300, hallsCount=1,
         f_water="да", f_ceremony="да", f_countryside="да", f_lodging="да", f_ownKitchen="да",
         phone="+7 495 755 39 43", email="banketfishpoint@gmail.com", site="https://fishpointrest.ru/",
         materials="https://drive.google.com/drive/folders/188zBL_ulAUPKVz8qiqil", priceDoc="https://fishpointrest.ru/banquet",
         slug="fish-point", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/fish-point/", address="Московская область, г.о. Подольск, д. Бережки", geo="55.386345, 37.618739", rating=5, reviewsCount=2586),
    dict(status="опубликовано", name="Усадьба Гребнево", type="усадьба", kicker="Усадьба · Щёлково",
         city="Щёлково", district="Подмосковье", highway="Щёлковское", kmFromMkad=25,
         capacityBanquet=350, capacityBuffet=600, hallsCount=9, checkFrom=7000, rentFrom=30000,
         f_water="да", f_ceremony="да", f_countryside="да", f_lodging="да", f_ownKitchen="да",
         phone="+7 915 437-11-82", email="banket@usadbagrebnevo.com", site="https://usadbagrebnevo.ru/banketnye-zaly/",
         materials="9 дисков по залам",
         slug="grebnevo", siteUrl="https://wed-secrets.ru/moskva/ploshchadki/grebnevo/", address="Московская область, г.о. Щёлково, усадьба Гребнево", geo="55.949301, 38.087608", rating=4.9, reviewsCount=55),
]

HALLS_RAW = [
    ("trubetskoy-hall", "Трубецкой Холл", "Гранд-шатёр", 740, 350, 600, [80000, 120000, 120000, 100000], [220000, 300000, 450000, 250000]),
    ("grand-lesnoy", "Grand «Лесной»", "Банкетный зал", 430, 280, 350, [200000, 300000, 400000, 200000], [350000, 700000, 800000, 400000]),
    ("karetny-dvor", "Каретный двор", "Банкетный зал", 260, 130, 170, [80000, 130000, 130000, 85000], [250000, 300000, 350000, 250000]),
    ("pushkin", "Пушкинъ", "Стеклянная оранжерея", 254, 130, 160, [150000, 200000, 300000, 200000], [300000, 350000, 450000, 300000]),
    ("lermontov", "Лермонтов", "Стеклянная оранжерея", 180, 90, 120, [120000, 200000, 250000, 180000], [250000, 300000, 380000, 300000]),
    ("barsky", "Барский", "Банкетный зал", 180, 80, 120, [100000, 150000, 150000, 100000], [250000, 380000, 400000, 250000]),
    ("shater-u-ozera", "Шатёр у озера", "Шатёр", 100, 50, 80, [60000, 80000, 80000, 60000], [150000, 200000, 250000, 150000]),
    ("limoncello", "Лимончелло", "Веранда", 80, 40, 60, [40000, 60000, 60000, 60000], [100000, 120000, 150000, 120000]),
    ("suvorov", "Суворов", "Лофт", 68, 30, 50, [30000, 50000, 50000, 30000], [110000, 150000, 200000, 110000]),
]
HALLS = [dict(venueName="Усадьба Гребнево", hallId=h[0], name=h[1], kicker=h[2], areaM2=h[3], banquet=h[4], buffet=h[5],
              rentWeekday=h[6][0], rentFri=h[6][1], rentSat=h[6][2], rentSun=h[6][3],
              depositWeekday=h[7][0], depositFri=h[7][1], depositSat=h[7][2], depositSun=h[7][3]) for h in HALLS_RAW]

SPECS = {}
for f in sorted(glob.glob("src/data/specialists/*/*.json")):
    s = json.load(open(f, encoding="utf-8"))
    SPECS.setdefault(s["categorySlug"], []).append(spec_row(s))
# Лешаковы — единственная карточка, живущая в TS, не в JSON
LESH = dict(status="опубликовано", name="Юлия и Леонид Лешаковы", who="пара", alsoCategories="Видеографы", tagline="Фото и видео одним комплектом",
            districts="все", mo="да", outsideMo="да", priceText="Зависит от даты, времени года и количества часов", experienceYears=15,
            style="Репортаж, Постановочный, Лайфстайл", serviceType="Полный день", materials="в репозитории", avatar="есть",
            slug="leshakovy", siteUrl="https://wed-secrets.ru/moskva/podryadchiki/fotografy/leshakovy/")
SPECS.setdefault("fotografy", []).insert(0, LESH)
SPECS.setdefault("videografy", []).insert(0, {**LESH, "style": "Репортаж", "notes": "Основная строка — на вкладке «Фотографы»"})


# ---------- сборка ----------
wb = Workbook()
wb.remove(wb.active)

def ref_range(name):
    col = list(REF.keys()).index(name) + 1
    L = get_column_letter(col)
    return f"'Справочники'!${L}$2:${L}${len(REF[name]) + 1}"

# Для Apps Script (registry-fix-validations.gs): списки по столбцам и какие
# столбцы многозначные — там выбор из списка ДОБАВЛЯЕТ значение, а не заменяет
GS_LISTS, GS_MULTI = {}, {}

def build_sheet(title, cols, rows, max_rows=300):
    ws = wb.create_sheet(title)
    GS_LISTS[title], GS_MULTI[title] = {}, []
    c = 1
    while c <= len(cols):
        g = cols[c - 1][2]
        e = c
        while e < len(cols) and cols[e][2] == g:
            e += 1
        fill = FILL_SITE if g == SITE else FILL_GROUP
        ws.cell(row=1, column=c, value=g).font = F_GROUP
        for cc in range(c, e + 1):
            ws.cell(row=1, column=cc).fill = fill
        if e > c:
            ws.merge_cells(start_row=1, start_column=c, end_row=1, end_column=e)
        c = e + 1
    ws.row_dimensions[3].height = 40

    for i, (key, label, group, kind, ex) in enumerate(cols, start=1):
        L = get_column_letter(i)
        ws.cell(row=2, column=i, value=key).font = F_KEY
        lc = ws.cell(row=3, column=i, value=label); lc.font = F_LABEL
        lc.fill = FILL_DIM if group == SITE else (FILL_REQ if ex.get("req") else FILL_LABEL)
        lc.alignment = WRAP; lc.border = BORDER
        note = ex.get("note", "")
        if kind.startswith("multi:"):
            note = (note + "\n" if note else "") + "Несколько через запятую: " + ", ".join(REF[kind[6:]])
        if note:
            lc.comment = Comment(note, "WED Secrets")
        w = {"long": 40, "text": 22, "url": 28, "num": 10, "int": 11, "yes": 9}.get(kind.split(":")[0], 18)
        if kind.startswith("multi:"): w = 28
        ws.column_dimensions[L].width = w
        rng = f"{L}4:{L}{max_rows}"
        dv = None
        if kind == "yes":
            dv = DataValidation(type="list", formula1='"да,нет"', allow_blank=True)
        elif kind.startswith("list:") or kind.startswith("multi:"):
            # Список — встроенный ("а,б,в"), а не ссылка на «Справочники»: Google Таблицы
            # при импорте xlsx делают ссылку на диапазон относительной, и у нижних строк
            # список уезжает — значения помечаются ошибкой. Ссылка остаётся только там,
            # где список не влезает в лимит xlsx (255 знаков).
            # multi — тот же список как подсказка; в Google Таблице включить «несколько вариантов».
            vals = REF[kind.split(":", 1)[1]]
            inline = ",".join(vals)
            formula = f'"{inline}"' if len(inline) <= 250 and not any("," in v for v in vals) else ref_range(kind.split(":", 1)[1])
            dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=False)
        elif kind == "int":
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
            dv.error = "Только число, без ₽ и пробелов"; dv.errorTitle = "Число"
        elif kind == "num":
            dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        if dv:
            ws.add_data_validation(dv); dv.add(rng)
        if kind == "yes":
            GS_LISTS[title][i] = ["да", "нет"]
        elif kind.startswith("list:") or kind.startswith("multi:"):
            GS_LISTS[title][i] = REF[kind.split(":", 1)[1]]
            if kind.startswith("multi:"):
                GS_MULTI[title].append(i)
        for r in range(4, 4 + max(len(rows), 1)):
            cell = ws.cell(row=r, column=i); cell.font = F_BODY
            cell.alignment = WRAP if kind in ("long", "text") or kind.startswith("multi:") else TOP
            if kind == "int": cell.number_format = "#,##0"

    keys = [c[0] for c in cols]
    for r, row in enumerate(rows, start=4):
        for key, val in row.items():
            if key in keys and val is not None:
                ws.cell(row=r, column=keys.index(key) + 1, value=val)
    ws.freeze_panes = "C4"  # шапка + статус + название
    ws.auto_filter.ref = f"A3:{get_column_letter(len(cols))}{max_rows}"
    return ws

ws = wb.create_sheet("Как заполнять")
ws.column_dimensions["A"].width = 110
GUIDE = [
    ("WED Secrets — реестр каталога", F_TITLE),
    ("Одна строка = одна карточка на сайте. Заказчик присылает текст и ссылки — раскладываем по столбцам сами.", F_BODY),
    ("", F_BODY),
    ("• «Площадки» — все площадки. Несколько залов (усадьба, комплекс) — площадка одной строкой здесь, залы по одному на строку на вкладке «Залы».", F_BODY),
    ("• Вкладки подрядчиков — по специализациям. Один подрядчик = одна строка в своём основном разделе; второй раздел — в столбце «Также в разделах».", F_BODY),
    ("• Жёлтая шапка — обязательно. Серая группа «Заполняет сайт» справа — трогать не надо (slug, страница, координаты, рейтинг: тянем сами).", F_BODY),
    ("• Строка 2 (серые ключи) нужна скрипту сборки сайта — не удалять.", F_BODY),
    ("• Числа — цифрами без ₽ и «от»: 120000. Пусто = цена по запросу. Да/нет — из списка; пусто = неизвестно.", F_BODY),
    ("• Столбцы с несколькими значениями (округа, стили) — через запятую; в Google Таблице можно включить «Разрешить несколько вариантов» в проверке данных — будут плашки.", F_BODY),
    ("• Соцсети — ссылки в одну ячейку, каждая с новой строки.", F_BODY),
    ("• Материалы — папка на диске с оригиналами фото (не сжатыми), подрядчику — портрет отдельным файлом.", F_BODY),
    ("• Статусы: черновик → в работе → на проверке → опубликовано; архив — снято с сайта.", F_BODY),
]
for r, (text, font) in enumerate(GUIDE, start=1):
    c = ws.cell(row=r, column=1, value=text); c.font = font; c.alignment = WRAP

build_sheet("Площадки", VENUE_COLS, VENUES)
build_sheet("Залы", HALL_COLS, HALLS)
for slug, label in CATEGORIES:
    specific, price_label = category_cols(slug)
    common = [(k, (f"Цена от, ₽ — {price_label.lower()}" if k == "priceFrom" and price_label else l), g, kind, ex)
              for k, l, g, kind, ex in SPEC_COMMON]
    build_sheet(label, common + specific + SPEC_CONTACTS, SPECS.get(slug, []))

ws = wb.create_sheet("Справочники")
for i, (name, vals) in enumerate(REF.items(), start=1):
    h = ws.cell(row=1, column=i, value=name); h.font = F_LABEL; h.fill = FILL_LABEL; h.alignment = WRAP
    ws.column_dimensions[get_column_letter(i)].width = max(14, min(26, max(len(v) for v in vals) + 2))
    for r, v in enumerate(vals, start=2):
        ws.cell(row=r, column=i, value=v).font = F_BODY
ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 30

wb.move_sheet("Справочники", offset=-(len(wb.sheetnames)-2))
wb.save(OUT)

# ---------- Apps Script для уже загруженной Google Таблицы ----------
GS = """// Сгенерировано scripts/registry-sheet.py — руками не править.
// 1) fixValidations — ставит выпадающие списки заново по карте «вкладка → столбец
//    → значения»: после импорта xlsx Google теряет ссылку на «Справочники».
// 2) onEdit — в многозначных столбцах (MULTI) выбор из списка ДОБАВЛЯЕТ значение
//    через запятую, повторный выбор убирает. Режим «несколько вариантов» через
//    API включить нельзя, поэтому так. Срабатывает сам, ничего запускать не надо.
const MAP = %s;
const MULTI = %s;
const FIRST_ROW = 4, LAST_ROW = 300;

function fixValidations() {
  const ss = SpreadsheetApp.getActiveSpreadsheet();
  let fixed = 0, missing = [];
  Object.keys(MAP).forEach((name) => {
    const sheet = ss.getSheetByName(name);
    if (!sheet) { missing.push(name); return; }
    if (sheet.getMaxRows() < LAST_ROW) sheet.insertRowsAfter(sheet.getMaxRows(), LAST_ROW - sheet.getMaxRows());
    Object.keys(MAP[name]).forEach((col) => {
      const rule = SpreadsheetApp.newDataValidation()
        .requireValueInList(MAP[name][col], true)
        .setAllowInvalid(true)
        .build();
      sheet.getRange(FIRST_ROW, Number(col), LAST_ROW - FIRST_ROW + 1, 1).setDataValidation(rule);
      fixed++;
    });
  });
  Logger.log('Готово: столбцов — ' + fixed + (missing.length ? '. Не найдены вкладки: ' + missing.join(', ') : ''));
}

function onEdit(e) {
  const range = e.range;
  const sheet = range.getSheet();
  const cols = MULTI[sheet.getName()];
  if (!cols || range.getNumRows() !== 1 || range.getNumColumns() !== 1) return;
  const col = range.getColumn();
  if (cols.indexOf(col) < 0 || range.getRow() < FIRST_ROW) return;
  const picked = String(e.value || '').trim();
  const old = String(e.oldValue || '').trim();
  if (!picked || !old || picked.indexOf(',') >= 0) return;   // пусто, первое значение или ручной ввод списка
  const list = old.split(',').map((v) => v.trim()).filter(String);
  const i = list.indexOf(picked);
  if (i >= 0) list.splice(i, 1); else list.push(picked);
  range.setValue(list.join(', '));
}
""" % (json.dumps(GS_LISTS, ensure_ascii=False, separators=(",", ":")),
       json.dumps(GS_MULTI, ensure_ascii=False, separators=(",", ":")))
open("scripts/registry-fix-validations.gs", "w", encoding="utf-8").write(GS)
print("gs:", sum(len(v) for v in GS_MULTI.values()), "многозначных столбцов")
print("saved", OUT, "sheets:", len(wb.sheetnames))
